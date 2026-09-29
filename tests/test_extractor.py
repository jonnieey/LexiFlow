"""Regression coverage for lexiflow.extractor (MetadataExtractor, fill_template).

This module was ported from LegatoFlow (src/lexiflow/extractor.py) but
had zero test coverage until now, despite being fully wired-in and
functional. Every test that instantiates MetadataExtractor mocks
lexiflow.extractor.OpenAI -- the real ~/.config/legatoflow/config on
this machine holds real API keys from prior manual use, and
MetadataExtractor.__init__ builds a real OpenAI client from
settings.OPENAI_API_KEY, so an unmocked instantiation would be able to
make real network calls with real credentials.
"""

import json
from pathlib import Path
from unittest.mock import MagicMock, patch

import pytest
from docx import Document

from lexiflow.extractor import MetadataExtractor, fill_template, get_ordinal_suffix


# -------- get_ordinal_suffix (pure) --------


@pytest.mark.parametrize(
    "num,expected",
    [
        (1, "1st"),
        (2, "2nd"),
        (3, "3rd"),
        (4, "4th"),
        (11, "11th"),
        (12, "12th"),
        (13, "13th"),
        (21, "21st"),
        (22, "22nd"),
        (23, "23rd"),
        (30, "30th"),
    ],
)
def test_get_ordinal_suffix(num, expected):
    assert get_ordinal_suffix(num) == expected


# -------- MetadataExtractor construction --------


@pytest.fixture
def extractor(monkeypatch):
    monkeypatch.setattr("lexiflow.extractor.settings.OPENAI_API_KEY", "sk-test")
    with patch("lexiflow.extractor.OpenAI") as mock_openai_cls:
        mock_openai_cls.return_value = MagicMock()
        yield MetadataExtractor()


def test_missing_ai_key_raises_clear_error(monkeypatch):
    monkeypatch.setattr("lexiflow.extractor.settings.OPENAI_API_KEY", "")
    monkeypatch.setattr(
        "lexiflow.extractor.settings.AI_API_KEY_ENV", "DEEPSEEK_API_KEY"
    )
    with patch("lexiflow.extractor.OpenAI") as mock_openai_cls:
        with pytest.raises(ValueError, match="DEEPSEEK_API_KEY is not set"):
            MetadataExtractor()
        mock_openai_cls.assert_not_called()


# -------- attorney name cleaning --------


def test_clean_attorney_name_strips_esq_variants(extractor):
    assert extractor._clean_attorney_name("Jane Doe, ESQUIRE") == "Jane Doe"
    assert extractor._clean_attorney_name("Jane Doe ESQ") == "Jane Doe"
    assert extractor._clean_attorney_name("Jane Doe") == "Jane Doe"


def test_clean_attorney_name_leaves_trailing_period_after_esq_dot(extractor):
    """Documents existing behavior, not a spec: the `\\b` word-boundary
    check after the optional `.` in `ESQ\\.?` doesn't match at end of
    string once the period is consumed, so the regex backtracks to
    leave the period behind. Pre-existing quirk, not fixed here."""
    assert extractor._clean_attorney_name("Jane Doe, ESQ.") == "Jane Doe."


def test_clean_attorney_name_handles_empty():
    ex = MetadataExtractor.__new__(MetadataExtractor)
    assert ex._clean_attorney_name("") == ""
    assert ex._clean_attorney_name(None) == ""


def test_get_attorney_last_name(extractor):
    assert extractor._get_attorney_last_name("Jane Q. Doe, ESQUIRE") == "Doe"
    assert extractor._get_attorney_last_name("") == ""


# -------- extract_notice (AI-mocked) --------


def test_extract_notice_happy_path(extractor, tmp_path):
    pdf_path = tmp_path / "notice.pdf"
    pdf_path.write_bytes(b"%PDF-1.4 fake")

    mock_message = MagicMock()
    mock_message.content = json.dumps(
        {"COURT_TYPE": "Circuit Court", "CASE_NUMBER": "2026-CV-001"}
    )
    mock_response = MagicMock()
    mock_response.choices = [MagicMock(message=mock_message)]
    extractor.client.chat.completions.create.return_value = mock_response

    with patch.object(extractor, "_extract_pdf_text", return_value="some notice text"):
        result = extractor.extract_notice(pdf_path)

    assert result == {"COURT_TYPE": "Circuit Court", "CASE_NUMBER": "2026-CV-001"}


def test_extract_notice_missing_file_returns_empty(extractor, tmp_path):
    result = extractor.extract_notice(tmp_path / "does-not-exist.pdf")
    assert result == {}


def test_extract_notice_ai_returns_invalid_json_returns_empty(extractor, tmp_path):
    pdf_path = tmp_path / "notice.pdf"
    pdf_path.write_bytes(b"%PDF-1.4 fake")

    mock_message = MagicMock()
    mock_message.content = "not valid json"
    mock_response = MagicMock()
    mock_response.choices = [MagicMock(message=mock_message)]
    extractor.client.chat.completions.create.return_value = mock_response

    with patch.object(extractor, "_extract_pdf_text", return_value="some notice text"):
        result = extractor.extract_notice(pdf_path)

    assert result == {}


# -------- extract_pbs branch selection --------


def test_extract_pbs_falls_back_to_rule_based_when_no_acroform_fields(extractor):
    with patch.object(extractor, "_extract_raw_acroform_fields", return_value={}):
        with patch.object(
            extractor, "_extract_pbs_rule_based", return_value={"JOB_NUMBER": "123"}
        ) as mock_rule_based:
            result = extractor.extract_pbs(Path("irrelevant.pdf"))

    mock_rule_based.assert_called_once()
    assert result == {"JOB_NUMBER": "123"}


def test_extract_pbs_uses_ai_result_when_valid(extractor):
    raw_fields = {"Witness 1": "Jane Doe"}
    ai_result = {"WITNESS_NAME": "Jane Doe", "attorneys": []}

    with patch.object(extractor, "_extract_raw_acroform_fields", return_value=raw_fields):
        with patch.object(extractor, "_ai_parse_acroform_fields", return_value=ai_result):
            with patch.object(
                extractor, "_extract_pbs_rule_based"
            ) as mock_rule_based:
                result = extractor.extract_pbs(Path("irrelevant.pdf"))

    mock_rule_based.assert_not_called()
    assert result.get("WITNESS_NAME") == "Jane Doe"


def test_extract_pbs_falls_back_when_ai_result_invalid(extractor):
    raw_fields = {"Witness 1": "Jane Doe"}
    ai_result = {}  # fails _validate_pbs_ai_output (no WITNESS_NAME/CASE_CAPTION)

    with patch.object(extractor, "_extract_raw_acroform_fields", return_value=raw_fields):
        with patch.object(extractor, "_ai_parse_acroform_fields", return_value=ai_result):
            with patch.object(
                extractor, "_extract_pbs_rule_based", return_value={"JOB_NUMBER": "456"}
            ) as mock_rule_based:
                result = extractor.extract_pbs(Path("irrelevant.pdf"))

    mock_rule_based.assert_called_once()
    assert result == {"JOB_NUMBER": "456"}


# -------- rule-based PBS field mapping (pure, via form_map) --------


def test_process_raw_pbs_map_maps_known_fields(extractor):
    extractor.form_map = {
        "Witness 1": "Jane Doe",
        "Job Date": "2026-03-15",
        "Start Time": "9:00",
        "AM": "On",
        "End Time": "5:00",
        "AM_2": "",
        "Court Reporter": "John Smith",
        "Job Number": "J-100",
        "Case Caption": "Doe v. Roe",
        "Yes": "On",
        "Combo Box7": "Reading waived",
    }

    extracted = extractor._process_raw_pbs_map()

    assert extracted["WITNESS_NAME"] == "Jane Doe"
    assert extracted["WITNESS_LAST_NAME"] == "Doe"
    assert extracted["DATE"] == "March 15, 2026"
    assert extracted["MONTH"] == "March"
    assert extracted["YEAR"] == "2026"
    assert extracted["START_TIME"] == "9:00 A.M."
    assert extracted["END_TIME"] == "5:00 P.M."
    assert extracted["COURT_REPORTER_NAME"] == "John Smith"
    assert extracted["JOB_NUMBER"] == "J-100"
    assert extracted["CASE_CAPTION"] == "Doe v. Roe"
    assert extracted["VIDEO_RECORDED"] == "Yes"
    assert extracted["WAIVER_OF_READING"] == "waived"


def test_process_raw_pbs_map_bad_date_falls_back_to_raw_string(extractor):
    extractor.form_map = {"Job Date": "not-a-date-at-all-xyz"}
    extracted = extractor._process_raw_pbs_map()
    assert extracted["RAW_DATE"] == "not-a-date-at-all-xyz"
    assert extracted["DATE"] == "not-a-date-at-all-xyz"


def test_process_raw_pbs_map_defaults_video_and_waiver_when_absent(extractor):
    extractor.form_map = {}
    extracted = extractor._process_raw_pbs_map()
    assert extracted["VIDEO_RECORDED"] == "No"
    assert extracted["WAIVER_OF_READING"] == "reserved"


# -------- attorney extraction fallback path --------


def test_extract_lawyers_groups_numbered_fields():
    ex = MetadataExtractor.__new__(MetadataExtractor)
    data = {
        "Attorney 1": "Jane Doe",
        "Form Name 1": "Doe Law Firm",
        "Attorney 2": "Yes",  # indicator value "yes" -> not a lawyer
        "Form Name 2": "Some Firm",
    }
    lawyers = ex.extract_lawyers(data, fields_to_extract=["Attorney", "Form Name"])
    assert len(lawyers) == 1
    assert lawyers[0]["Attorney"] == "Jane Doe"


def test_format_new_attorneys_builds_blocks():
    ex = MetadataExtractor.__new__(MetadataExtractor)
    lawyers = [{"Attorney": "Jane Doe, ESQUIRE", "Form Name": "Doe Law Firm"}]
    results = ex.format_new_attorneys(lawyers)
    assert results["TAKING_ATTORNEY1"] == "Jane Doe"
    assert results["TAKING_ATTORNEY1_LAST_NAME"] == "Doe"
    assert "ATTORNIES_BLOCK" in results


# -------- case-variant generation --------


def test_apply_case_variants_generates_upper_title_lower():
    ex = MetadataExtractor.__new__(MetadataExtractor)
    extracted = {
        "WITNESS_NAME": "jane doe",
        "END_TIME": "5:00 P.M.",
        "WAIVER_OF_READING": "WAIVED",
    }
    result = ex._apply_case_variants(extracted)
    assert result["WITNESS_NAME_UPPER"] == "JANE DOE"
    assert result["WITNESS_NAME_TITLE"] == "Jane Doe"
    assert result["END_TIME_LOWER"] == "5:00 p.m."
    assert result["WAIVER_OF_READING_LOWER"] == "waived"
    # original keys preserved
    assert result["WITNESS_NAME"] == "jane doe"


# -------- extract_all merge --------


def test_extract_all_merges_and_sorts(extractor):
    with patch.object(extractor, "extract_notice", return_value={"B_KEY": "1"}):
        with patch.object(extractor, "extract_pbs", return_value={"A_KEY": "2"}):
            result = extractor.extract_all(Path("notice.pdf"), Path("pbs.pdf"))

    assert list(result.keys()) == ["A_KEY", "B_KEY"]


def test_extract_all_notice_only(extractor):
    with patch.object(
        extractor, "extract_notice", return_value={"B_KEY": "1"}
    ) as mock_notice:
        with patch.object(extractor, "extract_pbs") as mock_pbs:
            result = extractor.extract_all(Path("notice.pdf"), None)

    mock_notice.assert_called_once_with(Path("notice.pdf"))
    mock_pbs.assert_not_called()
    assert result == {"B_KEY": "1"}


def test_extract_all_pbs_only(extractor):
    with patch.object(extractor, "extract_notice") as mock_notice:
        with patch.object(
            extractor, "extract_pbs", return_value={"A_KEY": "2"}
        ) as mock_pbs:
            result = extractor.extract_all(None, Path("pbs.pdf"))

    mock_notice.assert_not_called()
    mock_pbs.assert_called_once_with(Path("pbs.pdf"))
    assert result == {"A_KEY": "2"}


def test_extract_all_raises_without_either(extractor):
    with pytest.raises(ValueError, match="At least one"):
        extractor.extract_all(None, None)


# -------- map_to_template --------


def test_map_to_template_wraps_keys_in_brackets(extractor):
    result = extractor.map_to_template({"WITNESS_NAME": "Jane Doe"})
    assert result == {"[WITNESS_NAME]": "Jane Doe"}


# -------- fill_template (real docx, no AI/network involved) --------


def test_fill_template_replaces_paragraph_placeholder(tmp_path):
    template_path = tmp_path / "template.docx"
    doc = Document()
    doc.add_paragraph("Witness: [WITNESS_NAME]")
    doc.save(template_path)

    output_path = fill_template(template_path, {"WITNESS_NAME": "Jane Doe"})

    assert output_path == template_path.parent / "template_filled.docx"
    assert output_path.exists()
    filled = Document(output_path)
    assert filled.paragraphs[0].text == "Witness: Jane Doe"


def test_fill_template_replaces_table_cell_placeholder(tmp_path):
    template_path = tmp_path / "template.docx"
    doc = Document()
    table = doc.add_table(rows=1, cols=1)
    table.rows[0].cells[0].text = "[JOB_NUMBER]"
    doc.save(template_path)

    output_path = fill_template(template_path, {"JOB_NUMBER": "J-100"})

    filled = Document(output_path)
    assert filled.tables[0].rows[0].cells[0].text == "J-100"


def test_fill_template_uses_explicit_output_path(tmp_path):
    template_path = tmp_path / "template.docx"
    doc = Document()
    doc.add_paragraph("[JOB_NUMBER]")
    doc.save(template_path)

    explicit_output = tmp_path / "custom" / "result.docx"
    output_path = fill_template(template_path, {"JOB_NUMBER": "J-1"}, explicit_output)

    assert output_path == explicit_output
    assert explicit_output.exists()
